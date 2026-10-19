"""
政策文档切分逻辑的测试——用构造出来的文本直接测 _chunk_page_text/build_chunks，
不依赖真实文件；用真实写出的 .docx/.pdf 文件测一遍"从磁盘读取"这条路径也是通的。
"""
import docx
import pytest
from pypdf import PdfWriter

from app.services.policy_chunker import (
    CHUNK_OVERLAP,
    CHUNK_SIZE,
    RawPage,
    UnsupportedDocumentTypeError,
    build_chunks,
    chunk_document,
    extract_pages,
)


def test_short_text_becomes_single_chunk():
    pages = [RawPage(page_number=1, text="这是一段很短的政策原文。")]
    chunks = build_chunks(pages)
    assert len(chunks) == 1
    assert chunks[0].content == "这是一段很短的政策原文。"
    assert chunks[0].page_number == 1
    assert chunks[0].chunk_index == 0


def test_long_paragraph_gets_split_with_overlap():
    long_para = "第" + "一二三四五六七八九十" * (CHUNK_SIZE // 10 + 20)  # 明显超过 CHUNK_SIZE
    pages = [RawPage(page_number=None, text=long_para)]
    chunks = build_chunks(pages)

    assert len(chunks) >= 2
    # 相邻 chunk 应该有重叠：前一个 chunk 的末尾一段应该出现在下一个 chunk 开头
    tail_of_first = chunks[0].content[-CHUNK_OVERLAP:]
    assert tail_of_first in chunks[1].content


def test_chunk_index_is_global_and_sequential_across_pages():
    pages = [RawPage(page_number=1, text="第一页内容"), RawPage(page_number=2, text="第二页内容")]
    chunks = build_chunks(pages)
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    assert chunks[0].page_number == 1
    assert chunks[-1].page_number == 2


def test_empty_page_produces_no_chunks():
    assert build_chunks([RawPage(page_number=1, text="")]) == []
    assert build_chunks([]) == []


def test_unsupported_extension_raises_clear_error(tmp_path):
    bad_file = tmp_path / "policy.doc"
    bad_file.write_text("旧版 doc 格式")
    with pytest.raises(UnsupportedDocumentTypeError):
        extract_pages(str(bad_file))


def test_extract_and_chunk_real_docx_file(tmp_path):
    docx_path = tmp_path / "policy.docx"
    document = docx.Document()
    document.add_paragraph("江苏省2026年普通高校招生考试报名办法")
    document.add_paragraph("第一条：符合下列条件的人员，可以申请报名参加普通高等学校招生全国统一考试。")
    document.save(str(docx_path))

    chunks = chunk_document(str(docx_path))
    assert len(chunks) >= 1
    assert "报名办法" in chunks[0].content
    assert chunks[0].page_number is None  # docx 没有页码概念


def test_extract_and_chunk_real_pdf_file(tmp_path):
    pdf_path = tmp_path / "policy.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)  # 空白页：pypdf 提取不到文字层是预期行为
    with open(pdf_path, "wb") as f:
        writer.write(f)

    # 空白页提取不到文字，chunk_document 应该返回空列表而不是报错——
    # 真实场景对应"扫描版 PDF 没有文字层"，调用方（policy_ingest.py）
    # 会把这种情况当成一个明确的、可读的错误提示给用户，而不是在这里崩溃。
    chunks = chunk_document(str(pdf_path))
    assert chunks == []


# ---- 标题感知切分 ----

_DOC = (
    "—4—\n"
    "四、专业考试\n（一）省统考\n"
    "省统考设音乐类、舞蹈类、表（导）演类和书法类等6类。报名时考生可报考其中1类或多类，志愿填报时只能选择其中1类填报。\n"
    "省统考采取全省统一命题、统一考试、统一评分的形式，考试科目和内容按有关通知执行，各类专业省统考合格线在公布成绩时划定。\n"
    "（二）校考\n"
    "经批准，少数高校和专业可组织校考，具体专业名单由校考高校公布。校考工作应在学校所在地组织，不得跨省设置考点，严格控制现场人数。\n"
)


def test_structured_chunks_split_by_heading_with_breadcrumb_and_no_page_marker():
    chunks = build_chunks([RawPage(page_number=4, text=_DOC)])
    assert len(chunks) == 2
    assert chunks[0].content.startswith("四、专业考试 > （一）省统考\n省统考设音乐类")
    assert chunks[1].content.startswith("四、专业考试 > （二）校考\n经批准")
    assert all("—4—" not in c.content for c in chunks)
    assert all(c.page_number == 4 for c in chunks)


def test_sentence_starting_with_numbering_is_not_a_heading():
    # 「（二）综合素质评价围绕……，」是带编号的整句正文（PDF 还常把它折行成半句），不能当标题劈开
    text = "三、评价\n（一）2025 年，我省普通高校招生统一考试实行“3+1+2”模\n式，包括语文、数学、外语三门统考科目以及三门学业水平选择性考试科目，考生首先选择物理或历史，再选两门，共计六门科目参加考试。\n"
    chunks = build_chunks([RawPage(page_number=3, text=text)])
    assert len(chunks) == 1
    assert "（一）2025 年，我省" in chunks[0].content and "式，包括语文" in chunks[0].content


def test_tiny_sections_are_merged_into_next_and_keep_their_headings():
    text = "二、报名\n（一）报名资格\n凡符合条件者均可报名。\n（二）文化考试报名\n" + "报考艺术类的考生均须参加全国统一考试，其报名办法与普通类考生相同，由全省统一组织，与高考报名同步进行。" * 2 + "\n"
    chunks = build_chunks([RawPage(page_number=2, text=text)])
    assert len(chunks) == 1
    assert "（一）报名资格" in chunks[0].content and "凡符合条件者均可报名" in chunks[0].content
    assert chunks[0].content.startswith("二、报名")


def test_section_continues_across_pages_and_uses_start_page():
    p1 = "五、录取\n" + "录取工作分批次依次进行，" * 10 + "\n本科提前批次志愿设置如下，" + "各设置20个院校专业组志愿，" * 5
    p2 = "每个院校专业组志愿设置4个专业志愿和1个专业服从调剂志愿。\n"
    chunks = build_chunks([RawPage(page_number=7, text=p1), RawPage(page_number=8, text=p2)])
    assert chunks[0].page_number == 7
    assert any("设置4个专业志愿" in c.content for c in chunks)


def test_structured_false_keeps_legacy_behaviour():
    legacy = build_chunks([RawPage(page_number=4, text=_DOC)], structured=False)
    assert "—4—" in legacy[0].content and not legacy[0].content.startswith("四、专业考试 >")


def test_docx_extension_but_legacy_doc_content_gives_clear_error(tmp_path):
    fake = tmp_path / "旧格式.docx"
    fake.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(UnsupportedDocumentTypeError, match="旧版 .doc"):
        extract_pages(str(fake))
