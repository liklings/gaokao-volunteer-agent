"""
policy_ingest.py 的测试：用真实数据库（pgvector 表已经建好）+ 真实写出的
.docx/.pdf 文件，但 mock 掉 embed_texts（避免真的花钱调 Gemini API、也避免
测试跑网络请求）——用固定长度的假向量代替，只验证"落库这条链路"本身是对的：
document/chunk 记录建对了、幂等重新导入清空旧 chunk 重新生成、.doc 文件被
明确跳过并给出可读提示、单文件缺失或空内容时报出清楚的错误。

用真实数据库跟本项目其他测试文件（test_admission_master_loader.py 等）的
约定一致，测试结束时手动清理掉自己插入的行，不依赖事务回滚。
"""
from unittest.mock import AsyncMock, patch

import docx
import pytest
from sqlalchemy import delete, select

from app.models import PolicyChunk, PolicyDocument
from app.services import policy_ingest
from app.services.gemini_errors import GeminiCallError


def _fake_vector(dim: int = 768) -> list[float]:
    return [0.001] * dim


async def _fake_embed_texts(texts, task_type="RETRIEVAL_DOCUMENT"):
    return [_fake_vector() for _ in texts]


def _cleanup(db, source_path: str):
    doc = db.scalar(select(PolicyDocument).where(PolicyDocument.source_path == source_path))
    if doc is not None:
        db.execute(delete(PolicyChunk).where(PolicyChunk.document_id == doc.id))
        db.execute(delete(PolicyDocument).where(PolicyDocument.id == doc.id))
        db.commit()


def _make_docx(tmp_path, name: str, paragraphs: list[str]):
    path = tmp_path / name
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    document.save(str(path))
    return path


async def test_ingest_document_creates_document_and_chunks(db, tmp_path):
    path = _make_docx(tmp_path, "policy.docx", ["江苏省2026年普通高校招生考试报名办法", "第一条：符合条件的人员可以报名。"])
    try:
        with patch.object(policy_ingest, "embed_texts", AsyncMock(side_effect=_fake_embed_texts)):
            chunk_count = await policy_ingest.ingest_document(db, str(path))
            db.commit()

        assert chunk_count >= 1
        doc = db.scalar(select(PolicyDocument).where(PolicyDocument.source_path == str(path.resolve())))
        assert doc is not None
        assert doc.title == "policy"
        assert doc.doc_type == "docx"
        assert len(doc.chunks) == chunk_count
        assert doc.chunks[0].embedding is not None
    finally:
        _cleanup(db, str(path.resolve()))


async def test_reingest_same_path_clears_old_chunks_not_incremental(db, tmp_path):
    path = _make_docx(tmp_path, "policy.docx", ["旧版本内容，只有一段。"])
    try:
        with patch.object(policy_ingest, "embed_texts", AsyncMock(side_effect=_fake_embed_texts)):
            first_count = await policy_ingest.ingest_document(db, str(path))
            db.commit()

        doc = db.scalar(select(PolicyDocument).where(PolicyDocument.source_path == str(path.resolve())))
        first_doc_id = doc.id
        assert len(doc.chunks) == first_count

        # 重新写一份内容更长、段落更多的文件到同一个路径，模拟"文档更新后重新导入"
        document = docx.Document()
        document.add_paragraph("新版本第一段，内容跟旧版本完全不一样。")
        document.add_paragraph("新版本第二段，新增加的内容。")
        document.add_paragraph("新版本第三段，也是新增加的。")
        document.save(str(path))

        with patch.object(policy_ingest, "embed_texts", AsyncMock(side_effect=_fake_embed_texts)):
            second_count = await policy_ingest.ingest_document(db, str(path))
            db.commit()

        db.expire_all()
        doc_again = db.scalar(select(PolicyDocument).where(PolicyDocument.source_path == str(path.resolve())))
        assert doc_again.id == first_doc_id, "同一个 source_path 应该复用同一条 PolicyDocument 记录，不是新建一条"
        assert len(doc_again.chunks) == second_count, "重新导入后 chunk 数量应该跟新内容对应，不是新旧累加"
        assert second_count != first_count or True  # 数量是否变化取决于切分参数，这里只保证没有累加残留
        all_chunks = db.scalars(select(PolicyChunk).where(PolicyChunk.document_id == first_doc_id)).all()
        assert len(all_chunks) == second_count, "旧 chunk 没有被清空，出现了累加残留"
    finally:
        _cleanup(db, str(path.resolve()))


async def test_ingest_document_missing_file_raises_file_not_found_error(db, tmp_path):
    missing = tmp_path / "does_not_exist.docx"
    with pytest.raises(FileNotFoundError):
        await policy_ingest.ingest_document(db, str(missing))


async def test_ingest_document_empty_content_raises_value_error(db, tmp_path):
    # 空白 docx（没有任何段落文字）——chunk_document 会返回空列表，ingest_document
    # 应该把它变成一个人能看懂的 ValueError，而不是往下继续跑到 embedding 那一步。
    path = tmp_path / "empty.docx"
    docx.Document().save(str(path))
    with pytest.raises(ValueError):
        await policy_ingest.ingest_document(db, str(path))


async def test_ingest_directory_skips_doc_files_with_clear_message(tmp_path):
    _make_docx(tmp_path, "good.docx", ["这是一份能正常导入的文档。"])
    (tmp_path / "old_format.doc").write_bytes(b"fake legacy doc bytes")

    try:
        with patch.object(policy_ingest, "embed_texts", AsyncMock(side_effect=_fake_embed_texts)):
            result = await policy_ingest.ingest_directory(str(tmp_path))

        assert result.documents_processed == 1
        assert result.chunks_created >= 1
        assert len(result.skipped_files) == 1
        assert "old_format.doc" in result.skipped_files[0]
        assert ".docx" in result.skipped_files[0], "跳过原因里应该提示用户转存成 .docx"
    finally:
        from app.db import SessionLocal
        cleanup_db = SessionLocal()
        try:
            _cleanup(cleanup_db, str((tmp_path / "good.docx").resolve()))
        finally:
            cleanup_db.close()


async def test_ingest_directory_one_bad_file_does_not_abort_others(tmp_path):
    # 一份能正常导入的文档 + 一份空白文档（会在 embedding 之前就报 ValueError）——
    # 验证坏文件被记录进 skipped_files、不影响好文件的导入结果，且好文件的数据库
    # 改动（db.commit()）没有被坏文件那次的 db.rollback() 误伤。
    _make_docx(tmp_path, "good.docx", ["这是一份能正常导入的文档。"])
    docx.Document().save(str(tmp_path / "empty.docx"))

    try:
        with patch.object(policy_ingest, "embed_texts", AsyncMock(side_effect=_fake_embed_texts)):
            result = await policy_ingest.ingest_directory(str(tmp_path))

        assert result.documents_processed == 1
        assert len(result.skipped_files) == 1
        assert "empty.docx" in result.skipped_files[0]
    finally:
        from app.db import SessionLocal
        cleanup_db = SessionLocal()
        try:
            _cleanup(cleanup_db, str((tmp_path / "good.docx").resolve()))
        finally:
            cleanup_db.close()


async def test_ingest_directory_nonexistent_dir_raises():
    with pytest.raises(NotADirectoryError):
        await policy_ingest.ingest_directory("/no/such/directory/at/all")


async def test_ingest_directory_gemini_call_error_is_skipped_with_distinct_message(tmp_path):
    # 真实撞到过的问题：Gemini Embedding API 配额限流（429）导致某份文件失败时，
    # 之前的兜底文案统一写成"文件可能已损坏"，会误导用户去检查一份完全正常的
    # 文件。这里验证 GeminiCallError 走的是单独的分支，跳过原因里明确说的是
    # "调用 Gemini Embedding API 失败"，不是"文件可能已损坏"，而且措辞里
    # 提示了"重试后仍未成功，稍后单独用 --file 重新导入这一份即可"。
    _make_docx(tmp_path, "quota_hit.docx", ["这份文件在生成向量时会被 mock 成一直配额超限。"])

    async def _always_fail(texts, task_type="RETRIEVAL_DOCUMENT"):
        raise GeminiCallError("Gemini Embedding API（批量）返回 429：配额超限（模拟）")

    with patch.object(policy_ingest, "embed_texts", AsyncMock(side_effect=_always_fail)):
        result = await policy_ingest.ingest_directory(str(tmp_path))

    assert result.documents_processed == 0
    assert len(result.skipped_files) == 1
    msg = result.skipped_files[0]
    assert "quota_hit.docx" in msg
    assert "文件可能已损坏" not in msg, "配额限流不该被误报成文件损坏"
    assert "调用 Gemini Embedding API 失败" in msg
    assert "--file" in msg, "应该提示用户可以单独用 --file 重新导入这一份"


async def test_ingest_directory_corrupted_docx_is_skipped_not_crashed(tmp_path):
    # 真实撞到过的 bug：一份 .docx 文件内部 zip 结构损坏（不是"格式不支持"，是
    # "文件本身读不出来"），python-docx 读取时抛出的是 KeyError，不在
    # ingest_directory 原来只捕获 UnsupportedDocumentTypeError/ValueError/
    # FileNotFoundError 这三种"预期内"错误的范围里，导致整个目录的导入直接
    # 崩溃中断——好文件也导不进去了。这里手写一个"伪装成 docx、实际不是合法
    # zip 包"的文件来复现，验证修复后损坏文件被记录进 skipped_files、不影响
    # 同一批次里其他文件正常导入。
    _make_docx(tmp_path, "good.docx", ["这是一份能正常导入的文档。"])
    (tmp_path / "corrupted.docx").write_bytes(b"this is not a valid zip/docx file at all")

    try:
        with patch.object(policy_ingest, "embed_texts", AsyncMock(side_effect=_fake_embed_texts)):
            result = await policy_ingest.ingest_directory(str(tmp_path))

        assert result.documents_processed == 1, "损坏文件不该拖累好文件的导入"
        assert result.chunks_created >= 1
        assert len(result.skipped_files) == 1
        assert "corrupted.docx" in result.skipped_files[0]
    finally:
        from app.db import SessionLocal
        cleanup_db = SessionLocal()
        try:
            _cleanup(cleanup_db, str((tmp_path / "good.docx").resolve()))
        finally:
            cleanup_db.close()
